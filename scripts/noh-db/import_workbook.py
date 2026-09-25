# -*- coding: utf-8 -*-
"""import_workbook — Excel原本を読み、raw JSONスナップショットへ写す。

これが唯一 Excel（Master）に触れる工程である。以降の工程は raw.json だけを読む。
列名・値は一切改変しない。各行に原シート名とExcel行番号を provenance として付す。

  実行: python scripts/noh-db/import_workbook.py

依存: openpyxl（pip install openpyxl）。exceljs は当該ブックを開けなかったため採用しない。
"""
import io
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

try:
    import openpyxl
except ImportError:
    sys.stderr.write("openpyxl が必要です: python -m pip install openpyxl\n")
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK = ROOT / "private-data" / "noh" / "noh_play_database_v0_12_249plays_L2_complete.xlsx"
RAW = ROOT / "data" / "generated" / "normalized" / "raw.json"

SOURCE_VERSION = "v0.12-249plays-L2"
SCHEMA_VERSION = "1"

SHEETS = [
    "PLAY", "PLAY_SCHOOL", "CHARACTER", "RELATIONSHIP", "LOCATION", "SOURCE",
    "TAG", "STORY_PATTERN", "PERFORMANCE", "INTERPRETATION",
    "TAG_MASTER", "LEVEL_MASTER", "TAG_ALIAS", "TAG_RELATION",
]


def cell(v):
    if v is None:
        return None
    if isinstance(v, str):
        return v
    return v  # 数値・真偽はそのまま（JSONで表現できる）


def main():
    if not WORKBOOK.exists():
        sys.stderr.write("✗ Excel原本が見つかりません: %s\n" % WORKBOOK)
        sys.stderr.write("  private-data/noh/ に配置してください（.gitignore 済み・非公開）。\n")
        sys.exit(1)

    print("\n▶ import_workbook — Excel → raw.json")
    wb = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)

    out = {
        "_meta": {
            "sourceVersion": SOURCE_VERSION,
            "schemaVersion": SCHEMA_VERSION,
            "generatedAt": datetime.now(timezone.utc).isoformat(),
        },
        "tables": {},
    }

    for name in SHEETS:
        if name not in wb.sheetnames:
            sys.stderr.write("✗ 必須シートがありません: %s\n" % name)
            sys.exit(1)
        ws = wb[name]
        it = ws.iter_rows(values_only=False)
        try:
            header_cells = next(it)
        except StopIteration:
            header_cells = []
        cols = [str(c.value) if c.value is not None else None for c in header_cells]

        rows = []
        for row in it:
            rec = {"__sheet": name, "__row": row[0].row if row else None}
            has_any = False
            for c, cel in zip(cols, row):
                if not c:
                    continue
                val = cell(cel.value)
                rec[c] = val
                if val is not None and str(val).strip() != "":
                    has_any = True
            if has_any:  # 完全空行は取り込まない（TAG_MASTER末尾の空行など）
                rows.append(rec)

        out["tables"][name] = {"columns": [c for c in cols if c], "rows": rows}
        print("  ✓ %-14s %4d 行 / %2d 列" % (name, len(rows), len([c for c in cols if c])))

    RAW.parent.mkdir(parents=True, exist_ok=True)
    io.open(RAW, "w", encoding="utf-8").write(json.dumps(out, ensure_ascii=False, indent=2) + "\n")
    print("  ✓ 書き出し: data/generated/normalized/raw.json")


if __name__ == "__main__":
    main()
